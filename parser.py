"""
Парсер книги покупок (Раздел 8) и книги продаж (Раздел 9)
из официальной выгрузки декларации по НДС (формат ФНС).

Формат стабилен (регламентирован приказом ФНС), но на всякий случай
парсер ищет строку заголовка по ключевым словам, а не жёстко по номеру строки.
"""

from __future__ import annotations
import openpyxl
from dataclasses import dataclass, field
from datetime import datetime
from collections import defaultdict


@dataclass
class Operation:
    kind: str          # "purchase" | "sale"
    row_no: str
    invoice_no: str
    invoice_date: str
    inn: str
    kpp: str
    name: str
    amount: float       # стоимость операции (с НДС), в рублях
    vat: float | None    # сумма НДС, если применимо


@dataclass
class CounterpartySummary:
    inn: str
    name: str
    kpp: str
    purchase_total: float = 0.0
    purchase_count: int = 0
    sale_total: float = 0.0
    sale_count: int = 0
    operations: list = field(default_factory=list)

    @property
    def total_turnover(self) -> float:
        return self.purchase_total + self.sale_total

    @property
    def op_count(self) -> int:
        return self.purchase_count + self.sale_count


class ParseError(Exception):
    """Понятная ошибка, которую можно показать пользователю в боте."""


def _find_header_row(ws, must_contain=("ИНН", "Наименование"), max_scan=30):
    """Ищем строку заголовка по наличию ключевых слов, не завязываясь на номер строки."""
    for row_idx in range(1, min(max_scan, ws.max_row) + 1):
        row_values = [str(c.value) for c in ws[row_idx] if c.value is not None]
        joined = " ".join(row_values)
        if all(keyword in joined for keyword in must_contain):
            return row_idx
    raise ParseError(
        "Не нашёл строку заголовка с колонками ИНН/Наименование. "
        "Похоже, формат файла отличается от ожидаемого."
    )


def _col_index(header_row, keyword, occurrence=1, contains_all=None):
    """Находит номер колонки (1-based) по вхождению ключевого слова.
    occurrence=2 берёт второе совпадение (например, второй столбец 'ИНН')."""
    found = 0
    for idx, cell in enumerate(header_row, start=1):
        val = str(cell.value) if cell.value is not None else ""
        if keyword in val:
            if contains_all and not all(k in val for k in contains_all):
                continue
            found += 1
            if found == occurrence:
                return idx
    raise ParseError(f"Не нашёл колонку '{keyword}' (occurrence={occurrence}) в заголовке книги.")


def _parse_amount(val) -> float:
    if val is None or val == "":
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    try:
        return float(str(val).replace(" ", "").replace(",", "."))
    except ValueError:
        return 0.0


def _parse_date(val) -> str:
    if isinstance(val, datetime):
        return val.strftime("%d.%m.%Y")
    return str(val) if val else ""


def parse_pokupki(path: str) -> list[Operation]:
    """Раздел 8 — книга покупок. Контрагент = продавец (стр. 130)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    header_idx = _find_header_row(ws)
    header_row = ws[header_idx]

    col_row_no = _col_index(header_row, "№", occurrence=1)
    col_invoice_no = _col_index(header_row, "№", occurrence=2)
    col_invoice_date = _col_index(header_row, "Дата", occurrence=1)
    col_inn = _col_index(header_row, "ИНН", occurrence=1, contains_all=["130"])
    col_kpp = _col_index(header_row, "КПП", occurrence=1, contains_all=["130"])
    col_name = _col_index(header_row, "Наименование", occurrence=1)
    col_amount = _col_index(header_row, "Стоимость покупок")
    col_vat = _col_index(header_row, "Сумма НДС")

    ops = []
    for row in ws.iter_rows(min_row=header_idx + 1, values_only=False):
        inn = row[col_inn - 1].value
        if not inn:
            continue
        ops.append(Operation(
            kind="purchase",
            row_no=str(row[col_row_no - 1].value or ""),
            invoice_no=str(row[col_invoice_no - 1].value or ""),
            invoice_date=_parse_date(row[col_invoice_date - 1].value),
            inn=str(inn).strip(),
            kpp=str(row[col_kpp - 1].value or "").strip(),
            name=str(row[col_name - 1].value or "").strip(),
            amount=_parse_amount(row[col_amount - 1].value),
            vat=_parse_amount(row[col_vat - 1].value),
        ))
    return ops


def parse_prodazhi(path: str) -> list[Operation]:
    """Раздел 9 — книга продаж. Контрагент = покупатель (стр. 100)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    header_idx = _find_header_row(ws)
    header_row = ws[header_idx]

    col_row_no = _col_index(header_row, "№", occurrence=1)
    col_invoice_no = _col_index(header_row, "№", occurrence=2)
    col_invoice_date = _col_index(header_row, "Дата", occurrence=1)
    col_inn = _col_index(header_row, "ИНН", occurrence=1, contains_all=["100"])
    col_kpp = _col_index(header_row, "КПП", occurrence=1, contains_all=["100"])
    col_name = _col_index(header_row, "Наименование", occurrence=1)
    col_amount = _col_index(header_row, "в руб")

    ops = []
    for row in ws.iter_rows(min_row=header_idx + 1, values_only=False):
        inn = row[col_inn - 1].value
        if not inn:
            continue
        ops.append(Operation(
            kind="sale",
            row_no=str(row[col_row_no - 1].value or ""),
            invoice_no=str(row[col_invoice_no - 1].value or ""),
            invoice_date=_parse_date(row[col_invoice_date - 1].value),
            inn=str(inn).strip(),
            kpp=str(row[col_kpp - 1].value or "").strip(),
            name=str(row[col_name - 1].value or "").strip(),
            amount=_parse_amount(row[col_amount - 1].value),
            vat=None,
        ))
    return ops


def summarize(operations: list[Operation]) -> dict[str, CounterpartySummary]:
    """Группировка операций по ИНН контрагента."""
    result: dict[str, CounterpartySummary] = {}
    for op in operations:
        s = result.setdefault(op.inn, CounterpartySummary(inn=op.inn, name=op.name, kpp=op.kpp))
        if op.kind == "purchase":
            s.purchase_total += op.amount
            s.purchase_count += 1
        else:
            s.sale_total += op.amount
            s.sale_count += 1
        s.operations.append(op)
    return result


if __name__ == "__main__":
    import sys

    pokupki_path = sys.argv[1]
    prodazhi_path = sys.argv[2]

    pokupki = parse_pokupki(pokupki_path)
    prodazhi = parse_prodazhi(prodazhi_path)

    print(f"Книга покупок: {len(pokupki)} операций")
    print(f"Книга продаж: {len(prodazhi)} операций")

    all_ops = pokupki + prodazhi
    summary = summarize(all_ops)
    print(f"\nВсего уникальных контрагентов: {len(summary)}")

    top = sorted(summary.values(), key=lambda s: s.total_turnover, reverse=True)[:10]
    print("\nТоп-10 контрагентов по обороту:")
    for s in top:
        print(f"  {s.name:45.45s} ИНН {s.inn:12s} | покупки {s.purchase_total:>12,.2f} ₽ ({s.purchase_count}) "
              f"| продажи {s.sale_total:>12,.2f} ₽ ({s.sale_count})")
